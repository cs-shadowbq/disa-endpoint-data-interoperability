"""Tests for extractors."""

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
import pytest

from xls_extractor.extractors import JSONExtractor, MarkdownExtractor
from xls_extractor.writers import JSONWriter, MarkdownWriter


class TestMarkdownExtractor:
    """Test MarkdownExtractor class."""

    def test_extracts_single_column_text(self, tmp_path):
        """Test extraction of single-column text data."""
        extractor = MarkdownExtractor(MarkdownWriter())

        sheet_data = [
            ["Introduction"],
            ["This is line 1"],
            ["This is line 2"],
            [""],
            ["This is line 3"],
        ]

        extractor.extract(sheet_data, "Documentation", tmp_path)

        output_file = tmp_path / "Documentation.md"
        assert output_file.exists()

        with open(output_file, "r", encoding="utf-8") as f:
            content = f.read()

        assert "# Documentation" in content
        assert "Introduction" in content
        assert "This is line 1" in content
        assert "This is line 2" in content
        assert "This is line 3" in content

    def test_can_handle_documentation_sheet(self):
        """Test can_handle method for documentation sheets."""
        extractor = MarkdownExtractor(MarkdownWriter())

        doc_sheet = [
            ["Title"],
            ["Text line 1"],
            ["Text line 2"],
        ]

        assert extractor.can_handle(doc_sheet) is True

    def test_cannot_handle_structured_sheet(self):
        """Test can_handle method for structured sheets."""
        extractor = MarkdownExtractor(MarkdownWriter())

        structured_sheet = [
            ["Name", "Age", "City"],
            ["Alice", 30, "NYC"],
            ["Bob", 25, "LA"],
        ]

        assert extractor.can_handle(structured_sheet) is False

    def test_extracts_with_formatting(self, tmp_path):
        """Test extraction with Excel formatting preserved."""
        extractor = MarkdownExtractor(MarkdownWriter())

        # Create a test workbook with formatting
        workbook = openpyxl.Workbook()
        sheet = workbook.active

        # Add bold and underlined text (should be ## heading)
        cell1 = sheet.cell(row=1, column=1, value="Bold and Underlined Title")
        cell1.font = Font(bold=True, underline="single")

        # Add bold text (should be ### heading)
        cell2 = sheet.cell(row=2, column=1, value="Bold Subtitle")
        cell2.font = Font(bold=True)

        # Add regular text
        cell3 = sheet.cell(row=3, column=1, value="Regular text content")

        # Add another bold and underlined
        cell4 = sheet.cell(row=4, column=1, value="Another Title")
        cell4.font = Font(bold=True, underline="single")

        extractor.extract_with_formatting(sheet, "FormattedDoc", tmp_path)

        output_file = tmp_path / "FormattedDoc.md"
        assert output_file.exists()

        with open(output_file, "r", encoding="utf-8") as f:
            content = f.read()

        assert "## Bold and Underlined Title" in content
        assert "### Bold Subtitle" in content
        assert "Regular text content" in content
        assert "## Another Title" in content

    def test_ignores_formatting_for_numbered_items(self, tmp_path):
        """Test that numbered list items get #### heading."""
        extractor = MarkdownExtractor(MarkdownWriter())

        workbook = openpyxl.Workbook()
        sheet = workbook.active

        # Add text starting with number and period (should be #### heading)
        cell1 = sheet.cell(row=1, column=1, value="1. First item")
        cell1.font = Font(bold=True)

        # Add text starting with number and period (should be #### heading)
        cell2 = sheet.cell(row=2, column=1, value="2. Second item")
        cell2.font = Font(bold=True, underline="single")

        # Add text starting with number but no period (should NOT be heading)
        cell3 = sheet.cell(row=3, column=1, value="123 Some text")
        cell3.font = Font(bold=True)

        # Add regular bold text (should be ### heading)
        cell4 = sheet.cell(row=4, column=1, value="Regular Bold Title")
        cell4.font = Font(bold=True)

        extractor.extract_with_formatting(sheet, "NumberedDoc", tmp_path)

        output_file = tmp_path / "NumberedDoc.md"
        assert output_file.exists()

        with open(output_file, "r", encoding="utf-8") as f:
            content = f.read()

        # Numbered items with period should be #### headings
        assert "#### 1. First item" in content
        assert "#### 2. Second item" in content

        # Number without period should NOT have heading markers
        assert "123 Some text" in content
        assert "### 123 Some text" not in content

        # Regular bold should still be ### heading
        assert "### Regular Bold Title" in content


class TestJSONExtractor:
    """Test JSONExtractor class."""

    def test_extracts_structured_data(self, tmp_path):
        """Test extraction of structured data with headers."""
        extractor = JSONExtractor(JSONWriter())

        sheet_data = [
            ["Name", "Age", "City"],
            ["Alice", 30, "NYC"],
            ["Bob", 25, "LA"],
            ["Charlie", 35, "Chicago"],
        ]

        extractor.extract(sheet_data, "Users", tmp_path)

        output_file = tmp_path / "Users.json"
        assert output_file.exists()

        with open(output_file, "r", encoding="utf-8") as f:
            data = json.load(f)

        assert "Users" in data
        assert len(data["Users"]) == 3
        assert data["Users"][0] == {"Name": "Alice", "Age": 30, "City": "NYC"}
        assert data["Users"][1] == {"Name": "Bob", "Age": 25, "City": "LA"}

    def test_cleans_header_names(self, tmp_path):
        """Test that header names are cleaned (spaces to underscores)."""
        extractor = JSONExtractor(JSONWriter())

        sheet_data = [
            ["First Name", "Last Name", "Email Address"],
            ["John", "Doe", "john@example.com"],
        ]

        extractor.extract(sheet_data, "Contacts", tmp_path)

        output_file = tmp_path / "Contacts.json"
        with open(output_file, "r", encoding="utf-8") as f:
            data = json.load(f)

        record = data["Contacts"][0]
        assert "First_Name" in record
        assert "Last_Name" in record
        assert "Email_Address" in record

    def test_skips_empty_rows(self, tmp_path):
        """Test that empty rows are skipped."""
        extractor = JSONExtractor(JSONWriter())

        sheet_data = [
            ["Name", "Value"],
            ["Item1", 100],
            ["", ""],  # Empty row
            ["Item2", 200],
            [None, None],  # Empty row
            ["Item3", 300],
        ]

        extractor.extract(sheet_data, "Data", tmp_path)

        output_file = tmp_path / "Data.json"
        with open(output_file, "r", encoding="utf-8") as f:
            data = json.load(f)

        # Should only have 3 records (empty rows skipped)
        assert len(data["Data"]) == 3

    def test_handles_empty_sheet(self, tmp_path):
        """Test handling of empty sheets."""
        extractor = JSONExtractor(JSONWriter())

        sheet_data = []

        extractor.extract(sheet_data, "Empty", tmp_path)

        output_file = tmp_path / "Empty.json"
        assert output_file.exists()

        with open(output_file, "r", encoding="utf-8") as f:
            data = json.load(f)

        assert data == {"Empty": []}

    def test_handles_single_row_sheet(self, tmp_path):
        """Test handling of sheets with only headers."""
        extractor = JSONExtractor(JSONWriter())

        sheet_data = [["Name", "Age", "City"]]

        extractor.extract(sheet_data, "HeaderOnly", tmp_path)

        output_file = tmp_path / "HeaderOnly.json"
        with open(output_file, "r", encoding="utf-8") as f:
            data = json.load(f)

        assert data == {"HeaderOnly": []}

    def test_can_handle_structured_data(self):
        """Test can_handle method for structured data."""
        extractor = JSONExtractor(JSONWriter())

        structured_sheet = [
            ["Name", "Age"],
            ["Alice", 30],
        ]

        assert extractor.can_handle(structured_sheet) is True

    def test_handles_mismatched_columns(self, tmp_path):
        """Test handling of rows with different column counts."""
        extractor = JSONExtractor(JSONWriter())

        sheet_data = [
            ["Name", "Age", "City"],
            ["Alice", 30],  # Missing city
            ["Bob", 25, "LA", "Extra"],  # Extra column
            ["Charlie"],  # Missing age and city
        ]

        extractor.extract(sheet_data, "Irregular", tmp_path)

        output_file = tmp_path / "Irregular.json"
        with open(output_file, "r", encoding="utf-8") as f:
            data = json.load(f)

        records = data["Irregular"]
        assert len(records) == 3

        # First record missing city
        assert records[0]["Name"] == "Alice"
        assert records[0]["Age"] == 30
        assert records[0].get("City") is None

        # Second record has all columns (extra ignored)
        assert records[1]["Name"] == "Bob"
        assert records[1]["Age"] == 25
        assert records[1]["City"] == "LA"

        # Third record missing age and city
        assert records[2]["Name"] == "Charlie"
        assert records[2].get("Age") is None
        assert records[2].get("City") is None
