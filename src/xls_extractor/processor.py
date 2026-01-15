"""Excel file processor."""

from pathlib import Path
from typing import List, Optional

import openpyxl
import openpyxl.utils
from openpyxl.worksheet.worksheet import Worksheet

from .extractors import JSONExtractor, MarkdownExtractor
from .writers import JSONWriter, MarkdownWriter


class ExcelProcessor:
    """Processes Excel files and extracts sheets."""

    def __init__(
        self,
        output_dir: Path,
        force_markdown_sheets: Optional[List[str]] = None,
        include_hidden: bool = False,
        extract_mode: str = "values",
    ):
        """
        Initialize Excel processor.

        Args:
            output_dir: Directory to write extracted data
            force_markdown_sheets: List of sheet names to force as markdown
            include_hidden: Include hidden columns in extraction
            extract_mode: Extract "values" or "formulas" from cells
        """
        self.output_dir = Path(output_dir)
        self.force_markdown_sheets = set(force_markdown_sheets or [])
        self.include_hidden = include_hidden
        self.extract_mode = extract_mode

        # Initialize extractors
        self.markdown_extractor = MarkdownExtractor(MarkdownWriter())
        self.json_extractor = JSONExtractor(JSONWriter())

    def process_file(self, filepath: Path) -> dict:
        """
        Process an Excel file and extract all sheets.

        Args:
            filepath: Path to Excel file

        Returns:
            Dictionary with processing results
        """
        # Load with data_only based on extract_mode
        data_only = self.extract_mode == "values"
        workbook = openpyxl.load_workbook(filepath, data_only=data_only, rich_text=True)
        results = {
            "file": str(filepath),
            "sheets_processed": [],
            "sheets_failed": [],
        }

        for sheet_name in workbook.sheetnames:
            try:
                sheet = workbook[sheet_name]
                self._process_sheet(sheet, sheet_name)
                results["sheets_processed"].append(sheet_name)
            except Exception as e:
                results["sheets_failed"].append({"sheet": sheet_name, "error": str(e)})

        workbook.close()
        return results

    def _process_sheet(self, sheet: Worksheet, sheet_name: str) -> None:
        """
        Process a single sheet.

        Args:
            sheet: Worksheet object
            sheet_name: Name of the sheet
        """
        # Convert sheet to 2D list (values only for detection)
        sheet_data = []
        for row in sheet.iter_rows(values_only=True):
            # Filter hidden columns if needed
            if self.include_hidden:
                sheet_data.append(list(row))
            else:
                visible_row = []
                for idx, cell_value in enumerate(row):
                    # Check if column is hidden (1-indexed)
                    col_letter = openpyxl.utils.get_column_letter(idx + 1)
                    if not sheet.column_dimensions[col_letter].hidden:
                        visible_row.append(cell_value)
                sheet_data.append(visible_row)

        # Determine extractor to use
        if sheet_name in self.force_markdown_sheets:
            # Forced to markdown - pass actual cells for formatting
            extractor = self.markdown_extractor
            self._extract_with_formatting(sheet, sheet_name, extractor)
        elif self.markdown_extractor.can_handle(sheet_data):
            # Auto-detected as documentation - pass actual cells for formatting
            extractor = self.markdown_extractor
            self._extract_with_formatting(sheet, sheet_name, extractor)
        else:
            # Default to JSON for structured data - pass cells for Symbol font detection
            self._extract_json_with_cells(sheet, sheet_name)

    def _extract_json_with_cells(self, sheet: Worksheet, sheet_name: str) -> None:
        """
        Extract JSON data with access to cell objects for font detection.

        Args:
            sheet: Worksheet object
            sheet_name: Name of the sheet
        """
        all_rows = list(sheet.iter_rows())

        if not all_rows or len(all_rows) < 2:
            data = {sheet_name: []}
            self.json_extractor.writer.write(data, self.output_dir, sheet_name)
            return

        # Get header row with cells
        header_row = all_rows[0]
        headers = []
        visible_col_indices = []

        for idx, cell in enumerate(header_row):
            col_letter = openpyxl.utils.get_column_letter(idx + 1)
            if self.include_hidden or not sheet.column_dimensions[col_letter].hidden:
                header_value = self.json_extractor._clean_header(cell.value)
                if header_value:
                    headers.append(header_value)
                    visible_col_indices.append(idx)

        # Get data rows
        records = []
        for row in all_rows[1:]:
            # Skip empty rows
            if not any(cell.value for cell in row):
                continue

            record = {}
            for header, col_idx in zip(headers, visible_col_indices):
                if col_idx < len(row):
                    cell = row[col_idx]
                    value = self.json_extractor._clean_value(cell.value, cell)
                    record[header] = value
                else:
                    record[header] = None

            if record:
                records.append(record)

        data = {sheet_name: records}
        self.json_extractor.writer.write(data, self.output_dir, sheet_name)

    def _extract_with_formatting(self, sheet: Worksheet, sheet_name: str, extractor: MarkdownExtractor) -> None:
        """
        Extract sheet with formatting information for markdown.

        Args:
            sheet: Worksheet object
            sheet_name: Name of the sheet
            extractor: Markdown extractor to use
        """
        # Pass cell objects instead of just values
        extractor.extract_with_formatting(sheet, sheet_name, self.output_dir)
