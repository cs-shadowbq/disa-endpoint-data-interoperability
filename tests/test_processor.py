"""Tests for Excel processor."""

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
import pytest

from xls_extractor.processor import ExcelProcessor


@pytest.fixture
def sample_excel_file(tmp_path):
    """Create a sample Excel file for testing."""
    filepath = tmp_path / "test.xlsx"
    workbook = openpyxl.Workbook()

    # Create a structured data sheet
    sheet1 = workbook.active
    sheet1.title = "Users"
    sheet1.append(["Name", "Age", "City"])
    sheet1.append(["Alice", 30, "NYC"])
    sheet1.append(["Bob", 25, "LA"])

    # Create a documentation sheet
    sheet2 = workbook.create_sheet("README")
    sheet2.append(["Introduction"])
    sheet2.append(["This is a documentation sheet."])
    sheet2.append(["It contains text only."])

    workbook.save(filepath)
    workbook.close()

    return filepath


class TestExcelProcessor:
    """Test ExcelProcessor class."""

    def test_processes_file_successfully(self, sample_excel_file, tmp_path):
        """Test successful processing of an Excel file."""
        output_dir = tmp_path / "output"
        processor = ExcelProcessor(output_dir)

        results = processor.process_file(sample_excel_file)

        assert results["file"] == str(sample_excel_file)
        assert len(results["sheets_processed"]) == 2
        assert "Users" in results["sheets_processed"]
        assert "README" in results["sheets_processed"]
        assert len(results["sheets_failed"]) == 0

    def test_creates_output_files(self, sample_excel_file, tmp_path):
        """Test that output files are created."""
        output_dir = tmp_path / "output"
        processor = ExcelProcessor(output_dir)

        processor.process_file(sample_excel_file)

        # Check JSON file for structured data
        users_file = output_dir / "Users.json"
        assert users_file.exists()

        with open(users_file, "r", encoding="utf-8") as f:
            data = json.load(f)
            assert "Users" in data
            assert len(data["Users"]) == 2

        # Check Markdown file for documentation
        readme_file = output_dir / "README.md"
        assert readme_file.exists()

        with open(readme_file, "r", encoding="utf-8") as f:
            content = f.read()
            assert "# README" in content
            assert "Introduction" in content

    def test_force_markdown_sheets(self, tmp_path):
        """Test forcing specific sheets to be extracted as Markdown."""
        # Create a file with structured data
        filepath = tmp_path / "test.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append(["Name", "Value"])
        sheet.append(["Item1", 100])
        workbook.save(filepath)
        workbook.close()

        # Force it to be extracted as Markdown
        output_dir = tmp_path / "output"
        processor = ExcelProcessor(output_dir, force_markdown_sheets=["Data"])

        processor.process_file(filepath)

        # Should create a Markdown file instead of JSON
        md_file = output_dir / "Data.md"
        json_file = output_dir / "Data.json"

        assert md_file.exists()
        assert not json_file.exists()

    def test_handles_nonexistent_file(self, tmp_path):
        """Test handling of nonexistent files."""
        output_dir = tmp_path / "output"
        processor = ExcelProcessor(output_dir)

        nonexistent_file = tmp_path / "nonexistent.xlsx"

        with pytest.raises(Exception):
            processor.process_file(nonexistent_file)

    def test_processes_multiple_sheets(self, tmp_path):
        """Test processing a file with multiple sheets."""
        filepath = tmp_path / "multi_sheet.xlsx"
        workbook = openpyxl.Workbook()

        # Create multiple sheets
        for i in range(5):
            sheet = workbook.create_sheet(f"Sheet{i}")
            sheet.append(["Column1", "Column2"])
            sheet.append([f"Value{i}", i * 10])

        workbook.save(filepath)
        workbook.close()

        output_dir = tmp_path / "output"
        processor = ExcelProcessor(output_dir)

        results = processor.process_file(filepath)

        # Should process all sheets (including the default one created by openpyxl)
        assert len(results["sheets_processed"]) >= 5

    def test_extract_with_hidden_columns(self, tmp_path):
        """Test that hidden columns are excluded by default and included with flag."""
        filepath = tmp_path / "hidden_cols.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Data"

        # Add data with column C hidden (more columns to ensure it's seen as structured data)
        sheet.append(["Name", "Age", "Secret", "City"])
        sheet.append(["Alice", 30, "PASSWORD1", "NYC"])
        sheet.append(["Bob", 25, "PASSWORD2", "LA"])
        sheet.append(["Charlie", 35, "PASSWORD3", "Chicago"])

        # Hide column C (Secret)
        sheet.column_dimensions['C'].hidden = True

        workbook.save(filepath)
        workbook.close()

        # Test without hidden columns (default)
        output_dir1 = tmp_path / "output_no_hidden"
        processor1 = ExcelProcessor(output_dir1, include_hidden=False)
        processor1.process_file(filepath)

        json_file1 = output_dir1 / "Data.json"
        with open(json_file1, "r") as f:
            data1 = json.load(f)

        # Should only have Name, Age, and City columns (Secret excluded)
        assert len(data1["Data"][0]) == 3
        assert "Name" in data1["Data"][0]
        assert "Age" in data1["Data"][0]
        assert "City" in data1["Data"][0]
        assert "Secret" not in data1["Data"][0]

        # Test with hidden columns included
        output_dir2 = tmp_path / "output_with_hidden"
        processor2 = ExcelProcessor(output_dir2, include_hidden=True)
        processor2.process_file(filepath)

        json_file2 = output_dir2 / "Data.json"
        with open(json_file2, "r") as f:
            data2 = json.load(f)

        # Should have all four columns
        assert len(data2["Data"][0]) == 4
        assert "Name" in data2["Data"][0]
        assert "Age" in data2["Data"][0]
        assert "Secret" in data2["Data"][0]
        assert "City" in data2["Data"][0]

    def test_extract_formulas(self, tmp_path):
        """Test extracting formulas instead of values."""
        filepath = tmp_path / "formulas.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Calculations"

        # Add data with formulas (enough rows to be seen as structured data)
        sheet["A1"] = "Number"
        sheet["B1"] = "Double"
        sheet["C1"] = "Triple"
        sheet["A2"] = 5
        sheet["B2"] = "=A2*2"
        sheet["C2"] = "=A2*3"
        sheet["A3"] = 10
        sheet["B3"] = "=A3*2"
        sheet["C3"] = "=A3*3"
        sheet["A4"] = 15
        sheet["B4"] = "=A4*2"
        sheet["C4"] = "=A4*3"

        workbook.save(filepath)
        workbook.close()

        # Test extracting formulas
        output_dir = tmp_path / "output_formulas"
        processor = ExcelProcessor(output_dir, extract_mode="formulas")
        processor.process_file(filepath)

        json_file = output_dir / "Calculations.json"
        with open(json_file, "r") as f:
            data = json.load(f)

        # Should have formulas, not values
        assert data["Calculations"][0]["Double"] == "=A2*2"
        assert data["Calculations"][1]["Double"] == "=A3*2"

    def test_symbol_font_conversion(self, tmp_path):
        """Test that Symbol font characters are converted to Unicode."""
        filepath = tmp_path / "symbol_font.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Symbols"

        # Add header
        sheet["A1"] = "Symbol"
        sheet["B1"] = "Description"
        sheet["C1"] = "Value"

        # Add data with Symbol font - using Greek letters that will be converted
        sheet["A2"] = "Copyright"
        sheet["B2"] = "ã"  # This will be converted to ©
        sheet["B2"].font = Font(name='Symbol')
        sheet["C2"] = "Legal"

        sheet["A3"] = "Company"
        sheet["B3"] = "Χορπορατιον"  # Greek letters to be converted
        sheet["B3"].font = Font(name='Symbol')
        sheet["C3"] = "Name"

        # Add enough rows to make it structured data
        sheet["A4"] = "Test1"
        sheet["B4"] = "Value1"
        sheet["C4"] = "Type1"

        sheet["A5"] = "Test2"
        sheet["B5"] = "Value2"
        sheet["C5"] = "Type2"

        workbook.save(filepath)
        workbook.close()

        # Extract
        output_dir = tmp_path / "output"
        processor = ExcelProcessor(output_dir)
        results = processor.process_file(filepath)

        # Debug: check processing results
        print(f"Processing results: {results}")

        # Debug: check what files were created
        import os
        if output_dir.exists():
            files = list(output_dir.iterdir())
            print(f"Created files: {files}")
        else:
            print("Output directory does not exist!")

        # Check for either JSON or Markdown file
        json_file = output_dir / "Symbols.json"
        md_file = output_dir / "Symbols.md"

        if json_file.exists():
            with open(json_file, "r") as f:
                data = json.load(f)

            # Check that Symbol font was converted
            assert data["Symbols"][0]["Description"] == "©"  # Copyright symbol
            assert data["Symbols"][1]["Description"] == "Corporation"  # Greek to Latin
        elif md_file.exists():
            # If it was detected as markdown, just check that the conversion happened
            with open(md_file, "r") as f:
                content = f.read()
            assert "©" in content  # Copyright symbol
            assert "Corporation" in content  # Greek to Latin
        else:
            raise FileNotFoundError(f"No output file found in {output_dir}")
